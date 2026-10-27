"""DocxParser — Multi-modal Microsoft Word (.docx) structural extractor.

Extracts paragraphs, headings, tables (in true document order), and XML
footnotes/endnotes from .docx manuscripts into the unified ``Document``
and ``Page`` data structures used by ``DocumentParser``, ``SectionSegmenter``,
``CitationExtractor``, and ``ReferenceListParser``.
"""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from integrity_checker.extraction.base import BasePDFParser, Document, Page
from integrity_checker.logging import get_logger

logger = get_logger(__name__)

_WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

_BIB_HEADING_RE = re.compile(
    r"^\s*(?:[IVXLCDM]+\.\s*|\d+(?:\.\d+)*\.?\s*)?"
    r"(?:references|bibliography|works\s+cited|literature\s+cited|"
    r"tài\s+liệu\s+tham\s+khảo|danh\s+mục\s+tài\s+liệu\s+tham\s+khảo)\s*$",
    re.IGNORECASE,
)

# IEEE / Journal trailing author biography detector after bibliography entries
_AUTHOR_BIO_RE = re.compile(
    r"^[A-ZÀ-Ỹ][A-ZÀ-Ỹa-zà-ỹ\s.\-]{2,45}\s+"
    r"(?:is\s+currently\s+(?:an?\s+)?(?:undergraduate|graduate|phd|master|professor|lecturer|researcher|student)|"
    r"received\s+the\s+(?:B\.?S\.?|M\.?S\.?|Ph\.?D\.?|B\.?E\.?|B\.?Tech)|"
    r"was\s+born\s+in\s+)",
)


class DocxParser(BasePDFParser):
    """Structural parser for Microsoft Word (.docx) manuscripts."""

    name = "docx"

    # Approximate characters per logical page when explicit page breaks are absent
    CHARS_PER_PAGE = 3200

    def parse(self, file_path: str | Path) -> Document:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"DOCX file not found: {path}")

        try:
            import docx  # type: ignore
        except ImportError as exc:
            raise RuntimeError("python-docx is required to parse .docx files") from exc

        doc = docx.Document(str(path))
        footnotes_map = self._extract_xml_notes(path, "word/footnotes.xml", "footnote")
        endnotes_map = self._extract_xml_notes(path, "word/endnotes.xml", "endnote")

        pages_lines: list[list[str]] = [[]]
        current_chars = 0
        current_page_notes: list[str] = []
        in_bibliography = False

        def _flush_page() -> None:
            nonlocal current_chars, current_page_notes
            if current_page_notes:
                pages_lines[-1].extend(current_page_notes)
                current_page_notes = []
            if pages_lines[-1]:
                pages_lines.append([])
                current_chars = 0

        body_elm = doc.element.body
        for child in body_elm.iterchildren():
            tag = child.tag
            if tag == f"{_WORD_NS}p":
                text, note_refs, has_page_break = self._extract_paragraph_xml(
                    child, footnotes_map, endnotes_map
                )
                if has_page_break and pages_lines[-1]:
                    _flush_page()

                stripped = text.strip()
                if not stripped:
                    continue

                # Check if this paragraph is the Bibliography / References heading
                if _BIB_HEADING_RE.match(stripped):
                    if pages_lines[-1]:
                        _flush_page()
                    in_bibliography = True
                    pages_lines[-1].append(stripped)
                    current_chars += len(stripped)
                    continue

                # Skip trailing IEEE author biography paragraphs at the end of REFERENCES
                if in_bibliography and _AUTHOR_BIO_RE.match(stripped):
                    logger.debug(
                        "Skipping trailing author biography paragraph in DOCX: {}",
                        stripped[:80],
                    )
                    continue

                pages_lines[-1].append(stripped)
                current_chars += len(stripped)

                for note_txt in note_refs:
                    if note_txt not in current_page_notes:
                        current_page_notes.append(note_txt)

                if not in_bibliography and current_chars >= self.CHARS_PER_PAGE:
                    _flush_page()

            elif tag == f"{_WORD_NS}tbl":
                table_lines = self._extract_table_xml(child)
                for tline in table_lines:
                    pages_lines[-1].append(tline)
                    current_chars += len(tline)
                if not in_bibliography and current_chars >= self.CHARS_PER_PAGE:
                    _flush_page()

        if current_page_notes:
            pages_lines[-1].extend(current_page_notes)

        # Append any unreferenced endnotes at the end of body if not already included
        all_joined = "\n".join("\n".join(lines) for lines in pages_lines)
        for eid, etext in sorted(endnotes_map.items()):
            marker = f"[{eid}] {etext}"
            if etext and etext not in all_joined:
                pages_lines[-1].append(marker)

        non_empty_pages = [
            "\n".join(lines).strip()
            for lines in pages_lines
            if "\n".join(lines).strip()
        ]
        if not non_empty_pages:
            return Document(
                file_path=str(path),
                num_pages=1,
                pages=[Page(page_num=1, text="")],
                parser_used=self.name,
            )

        pages = [
            Page(page_num=idx + 1, text=page_text)
            for idx, page_text in enumerate(non_empty_pages)
        ]
        return Document(
            file_path=str(path),
            num_pages=len(pages),
            pages=pages,
            parser_used=self.name,
        )

    @staticmethod
    def _extract_paragraph_xml(
        p_elm,
        footnotes_map: dict[int, str],
        endnotes_map: dict[int, str],
    ) -> tuple[str, list[str], bool]:
        """Extract text, referenced footnotes/endnotes, and page break flag from a <w:p> element."""
        parts: list[str] = []
        note_texts: list[str] = []
        has_page_break = False

        for elem in p_elm.iter():
            etag = elem.tag
            if etag == f"{_WORD_NS}t" and elem.text:
                parts.append(elem.text)
            elif etag == f"{_WORD_NS}tab":
                parts.append(" ")
            elif etag in (f"{_WORD_NS}br", f"{_WORD_NS}cr"):
                br_type = elem.attrib.get(f"{_WORD_NS}type", "")
                if br_type == "page":
                    has_page_break = True
                else:
                    parts.append("\n")
            elif etag == f"{_WORD_NS}lastRenderedPageBreak":
                has_page_break = True
            elif etag == f"{_WORD_NS}footnoteReference":
                raw_id = elem.attrib.get(f"{_WORD_NS}id")
                if raw_id and raw_id.lstrip("-").isdigit():
                    fid = int(raw_id)
                    if fid > 0 and fid in footnotes_map:
                        note_texts.append(f"{fid}. {footnotes_map[fid]}")
            elif etag == f"{_WORD_NS}endnoteReference":
                raw_id = elem.attrib.get(f"{_WORD_NS}id")
                if raw_id and raw_id.lstrip("-").isdigit():
                    eid = int(raw_id)
                    if eid > 0 and eid in endnotes_map:
                        note_texts.append(f"{eid}. {endnotes_map[eid]}")

        text = "".join(parts)
        return text, note_texts, has_page_break

    @staticmethod
    def _extract_table_xml(tbl_elm) -> list[str]:
        """Extract distinct row texts from a <w:tbl> element."""
        lines: list[str] = []
        for tr in tbl_elm.findall(f".//{_WORD_NS}tr"):
            row_cells: list[str] = []
            for tc in tr.findall(f"./{_WORD_NS}tc"):
                cell_texts = [
                    t.text for t in tc.findall(f".//{_WORD_NS}t") if t.text
                ]
                cell_str = " ".join("".join(cell_texts).split())
                if cell_str and (not row_cells or row_cells[-1] != cell_str):
                    row_cells.append(cell_str)
            if row_cells:
                lines.append(" | ".join(row_cells))
        return lines

    @staticmethod
    def _extract_xml_notes(
        docx_path: Path, member_name: str, note_tag: str
    ) -> dict[int, str]:
        """Extract footnotes or endnotes from word/footnotes.xml or word/endnotes.xml."""
        notes: dict[int, str] = {}
        try:
            with zipfile.ZipFile(docx_path, "r") as zf:
                if member_name not in zf.namelist():
                    return notes
                xml_bytes = zf.read(member_name)
            root = ET.fromstring(xml_bytes)
            for note_el in root.findall(f".//{_WORD_NS}{note_tag}"):
                raw_id = note_el.attrib.get(f"{_WORD_NS}id")
                note_type = note_el.attrib.get(f"{_WORD_NS}type", "")
                if note_type in ("separator", "continuationSeparator"):
                    continue
                if not raw_id or not raw_id.lstrip("-").isdigit():
                    continue
                nid = int(raw_id)
                if nid <= 0:
                    continue
                texts = [
                    t.text for t in note_el.findall(f".//{_WORD_NS}t") if t.text
                ]
                cleaned = " ".join("".join(texts).split())
                if cleaned:
                    notes[nid] = cleaned
        except Exception as exc:
            logger.debug("Could not parse {} from {}: {}", member_name, docx_path, exc)
        return notes
