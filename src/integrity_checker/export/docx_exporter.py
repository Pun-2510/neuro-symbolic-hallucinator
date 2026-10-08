"""DOCX export functionality for analysis reports."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def export_report_to_docx(
    filename: str,
    cis_score: float,
    num_citations: int,
    num_verified: int,
    num_metadata_error: int,
    num_suspected: int,
    num_unresolved: int,
    num_resource: int,
    verdicts: list[dict],
    generated_at: str | None = None,
) -> bytes:
    """Generate a DOCX report from analysis results.

    Args:
        filename: Name of the analyzed essay
        cis_score: Citation Integrity Score (0-100)
        num_citations: Total number of citations
        num_verified: Number of verified citations
        num_metadata_error: Number with metadata errors
        num_suspected: Number of suspected hallucinations
        num_unresolved: Number of unresolved citations
        num_resource: Number of URL/Resource citations
        verdicts: List of verdict dictionaries
        generated_at: Timestamp string

    Returns:
        bytes: DOCX file content
    """
    doc = Document()

    # Title
    title = doc.add_heading("Essay Integrity Check Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata
    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.add_run("Essay: ").bold = True
    meta.add_run(filename)
    meta = doc.add_paragraph()
    meta.add_run("Generated: ").bold = True
    meta.add_run(generated_at or datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"))
    meta = doc.add_paragraph()
    meta.add_run("System: ").bold = True
    meta.add_run("Essay Integrity Checker (DATN)")

    # CIS Score Section
    doc.add_heading("Citation Integrity Score (CIS)", level=1)

    score_para = doc.add_paragraph()
    score_run = score_para.add_run(f"{cis_score:.2f}%")
    score_run.font.size = Pt(36)
    score_run.font.bold = True

    # Determine score color
    if cis_score >= 90:
        score_run.font.color.rgb = RGBColor(34, 197, 94)  # Green
    elif cis_score >= 70:
        score_run.font.color.rgb = RGBColor(234, 179, 8)  # Yellow
    else:
        score_run.font.color.rgb = RGBColor(239, 68, 68)  # Red

    doc.add_paragraph("(Score represents percentage of citations that passed integrity verification)")

    # Summary Statistics
    doc.add_heading("Summary Statistics", level=1)

    stats_table = doc.add_table(rows=8, cols=2)
    stats_table.style = "Table Grid"

    stats_data = [
        ("Total Citations", str(num_citations)),
        ("Verified", f"{num_verified} ({num_verified * 100 / max(num_citations, 1):.1f}%)"),
        ("Metadata Error", f"{num_metadata_error} ({num_metadata_error * 100 / max(num_citations, 1):.1f}%)"),
        ("Suspected Hallucination", f"{num_suspected} ({num_suspected * 100 / max(num_citations, 1):.1f}%)"),
        ("Unresolved", f"{num_unresolved} ({num_unresolved * 100 / max(num_citations, 1):.1f}%)"),
        ("URL/Resource", f"{num_resource} ({num_resource * 100 / max(num_citations, 1):.1f}%)"),
        ("", ""),
        ("Citation Integrity Score", f"{cis_score:.2f}%"),
    ]

    for i, (label, value) in enumerate(stats_data):
        row = stats_table.rows[i]
        row.cells[0].text = label
        row.cells[1].text = value
        if label == "Citation Integrity Score":
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True

    # Verdicts Detail Section
    doc.add_heading("Citation Verdicts", level=1)

    if verdicts:
        # Verdicts table
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"

        # Header row
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Citation"
        hdr_cells[1].text = "Type"
        hdr_cells[2].text = "Label"
        hdr_cells[3].text = "Confidence"

        for cell in hdr_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        # Data rows
        for v in verdicts:
            row_cells = table.add_row().cells
            row_cells[0].text = v.get("citation_raw", "")[:100] + ("..." if len(v.get("citation_raw", "")) > 100 else "")
            row_cells[1].text = v.get("citation_type", "unknown")
            row_cells[2].text = v.get("label", "UNKNOWN")
            row_cells[3].text = f"{v.get('confidence', 0):.2f}"

    else:
        doc.add_paragraph("No citation verdicts available.")

    # Legend
    doc.add_heading("Label Legend", level=1)

    legend_table = doc.add_table(rows=6, cols=2)
    legend_table.style = "Table Grid"

    legend_data = [
        ("VERIFIED", "Citation verified successfully through academic databases"),
        ("METADATA_ERROR", "Citation found but metadata (author, title, year) doesn't match"),
        ("SUSPECTED_HALLUCINATION", "Citation appears to be fabricated or non-existent"),
        ("UNRESOLVED", "Insufficient evidence to determine citation validity"),
        ("RESOURCE", "Citation is a URL/web resource (not academic paper)"),
        ("", ""),
    ]

    for i, (label, desc) in enumerate(legend_data):
        if i >= len(legend_table.rows):
            break
        row = legend_table.rows[i]
        row.cells[0].text = label
        row.cells[1].text = desc
        if label == "VERIFIED":
            row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(34, 197, 94)
        elif label == "METADATA_ERROR":
            row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(234, 179, 8)
        elif label == "SUSPECTED_HALLUCINATION":
            row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(239, 68, 68)

    # Disclaimer
    doc.add_paragraph()
    disclaimer = doc.add_paragraph()
    disclaimer_run = disclaimer.add_run(
        "DISCLAIMER: This report is a decision-support tool only. "
        "Final judgment on citation integrity requires expert human review. "
        "The system may produce false positives or false negatives."
    )
    disclaimer_run.font.size = Pt(9)
    disclaimer_run.font.italic = True

    # Save to bytes
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
