"""Unit tests for all 4 P3 Killer Features:
1. Live Citation Sandbox & Synthetic Academic PDF Generator from ACL Database
2. Auto-Healed Word (.docx) Export with Before -> After Change Log
3. Revision Diff / Version Progress Tracker (v1 -> v2)
4. Scanned / Image-Only PDF Diagnostic Guardrail (SCANNED_PDF_DETECTED)
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import fitz
from docx import Document as WordDocument

from integrity_checker.api.routes.essays import (
    _format_in_text_marker_apa,
    _format_sandbox_ref_entry,
    _synthesize_academic_manuscript_pdf,
)
from integrity_checker.api.routes.report import (
    _compute_revision_diff,
    _extract_document_diagnostics,
    _normalize_manuscript_stem,
)
from integrity_checker.database import LocalDatabase
from integrity_checker.export.docx_exporter import export_auto_healed_docx
from integrity_checker.extraction import DocumentParser, MuPdfParser


def test_acl_random_sampler_and_synthetic_pdf_generation(tmp_path: Path) -> None:
    """Feature #1: Sample 12 papers from ACL DB, mutate metadata & linking, synthesize full PDF, and parse In-Text <-> Refs."""
    db_file = tmp_path / "test_acl.db"
    local_db = LocalDatabase(db_path=db_file)
    sampled = local_db.sample_random_papers(count=12)
    assert len(sampled) == 12

    items = []
    for idx, p in enumerate(sampled, start=1):
        items.append(
            {
                "index": idx,
                "paper_id": p.id,
                "title": p.title,
                "authors": p.authors,
                "authors_text": "; ".join(p.authors),
                "year": str(p.year or 2020),
                "venue": p.venue or "Proceedings of ACL",
                "doi": p.doi or "",
                "cite_in_text": True,
                "include_in_references": True,
                "in_text_year_override": "",
                "in_text_context": "",
            }
        )

    # Inject Layer-1 and Layer-2 test mutations
    items[1]["year"] = "2024"  # Metadata year error
    items[2]["in_text_year_override"] = "2011"  # In-text <-> Ref year mismatch
    items[3]["cite_in_text"] = False  # Uncited reference
    items[4]["include_in_references"] = False  # Missing reference (orphan in-text)

    pdf_bytes, num_pages, in_text_markers, ref_lines = _synthesize_academic_manuscript_pdf(
        title="Empirical Evaluation of Neural Language Representations in Computational Linguistics",
        author_name="CITRA Defense Sandbox",
        style="apa",
        items=items,
    )
    assert pdf_bytes.startswith(b"%PDF-")
    assert num_pages >= 1
    assert len(in_text_markers) == 11  # 12 - 1 uncited
    assert len(ref_lines) == 11  # 12 - 1 orphan in-text

    pdf_path = tmp_path / "synthesized_manuscript.pdf"
    pdf_path.write_bytes(pdf_bytes)

    parsed = DocumentParser().parse(str(pdf_path))
    assert parsed.document.metadata.get("is_scanned_pdf") is False
    assert len(parsed.body_citations) >= 10
    assert len(parsed.references) == 11


def test_auto_healed_docx_export_and_in_place_remediation(tmp_path: Path) -> None:
    """Feature #2: Generate Auto-Healed .docx with Before -> After Change Log and heal in-place .docx."""
    orig_doc = WordDocument()
    orig_doc.add_heading("Research Paper", level=1)
    orig_doc.add_paragraph("References")
    orig_doc.add_paragraph(
        "Vaswani, A., Shazeer, N., & Parmar, N. (2019). Attention Is All You Need. NeurIPS."
    )
    orig_doc.add_paragraph(
        "Sterling, J. (2025). Quantum Hyper-Graph Prompting for Imaginary AI. Fake Journal."
    )
    orig_docx_path = tmp_path / "draft_v1.docx"
    orig_doc.save(str(orig_docx_path))

    verdicts = [
        {
            "citation_raw": "Vaswani, A., Shazeer, N., & Parmar, N. (2019). Attention Is All You Need. NeurIPS.",
            "label": "metadata_error",
            "mismatched_fields": ["year"],
            "matched_sources": [
                {
                    "source_name": "acl",
                    "title": "Attention Is All You Need",
                    "authors": ["Ashish Vaswani", "Noam Shazeer", "Niki Parmar"],
                    "year": 2017,
                    "venue": "Advances in Neural Information Processing Systems (NeurIPS)",
                    "doi": "10.48550/arXiv.1706.03762",
                }
            ],
        },
        {
            "citation_raw": "Sterling, J. (2025). Quantum Hyper-Graph Prompting for Imaginary AI. Fake Journal.",
            "label": "suspected_hallucination",
            "mismatched_fields": [],
            "matched_sources": [],
        },
    ]

    healed_bytes = export_auto_healed_docx(
        essay_filename="draft_v1.docx",
        verdicts=verdicts,
        style="apa",
        original_docx_path=str(orig_docx_path),
    )
    assert len(healed_bytes) > 1000

    reopened = WordDocument(io.BytesIO(healed_bytes))
    full_text = "\n".join(p.text for p in reopened.paragraphs)
    assert "Auto-Healed by CITRA" in full_text
    assert "2017" in full_text
    assert "HALLUCINATION" in full_text
    assert len(reopened.tables) >= 1  # Before -> After Change Log table


def test_revision_diff_computation() -> None:
    """Feature #3: Verify manuscript stem normalization and v1 -> v2 revision diff calculation."""
    assert _normalize_manuscript_stem("Thesis_Manuscript_v1.pdf") == "thesis_manuscript"
    assert _normalize_manuscript_stem("Thesis_Manuscript_v2_fixed.docx") == "thesis_manuscript"

    class DummyVerdict:
        def __init__(self, raw: str, label: str) -> None:
            self.citation_raw = raw
            self.label = label

    class DummyEssay:
        def __init__(self, eid: int, filename: str, cis_score: float, user_id: int = 1) -> None:
            self.id = eid
            self.filename = filename
            self.user_id = user_id
            self.uploaded_at = None
            self.cis_json = json.dumps({"score": cis_score})

    class DummyRepo:
        def __init__(self, e1: DummyEssay, e2: DummyEssay, v1: list, v2: list) -> None:
            self._essays = [e2, e1]
            self._verdicts = {e1.id: v1, e2.id: v2}

        def get_user_essays(self, user_id: int) -> list:
            return self._essays

        def get_verdicts(self, essay_id: int) -> list:
            return self._verdicts.get(essay_id, [])

    e_v1 = DummyEssay(10, "KLTN_Draft_v1.pdf", 62.5)
    e_v2 = DummyEssay(11, "KLTN_Draft_v2_fixed.pdf", 94.0)
    v1_list = [
        DummyVerdict("Vaswani et al. (2019). Attention Is All You Need. 10.48550/arXiv.1706.03762", "metadata_error"),
        DummyVerdict("Fake Author (2025). Hallucinated Paper.", "suspected_hallucination"),
    ]
    v2_list = [
        DummyVerdict("Vaswani et al. (2017). Attention Is All You Need. 10.48550/arXiv.1706.03762", "verified"),
    ]
    repo = DummyRepo(e_v1, e_v2, v1_list, v2_list)

    diff = _compute_revision_diff(repo, e_v2, v2_list)  # type: ignore[arg-type]
    assert diff is not None
    assert diff["has_previous_revision"] is True
    assert diff["previous_essay_id"] == 10
    assert diff["cis_before"] == 62.5
    assert diff["cis_after"] == 94.0
    assert diff["cis_delta"] == 31.5
    assert diff["resolved_issues_count"] == 2


def test_scanned_pdf_guardrail_detection(tmp_path: Path) -> None:
    """Feature #4: Detect image-only/scanned PDFs vs digital text PDFs."""
    scanned_path = tmp_path / "scanned_image_only.pdf"
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    # Draw only a rectangle graphic, no extractable text layer
    page.draw_rect(fitz.Rect(50, 50, 500, 750), color=(0.2, 0.2, 0.2))
    doc.save(str(scanned_path))
    doc.close()

    mupdf_doc = MuPdfParser().parse(scanned_path)
    assert mupdf_doc.metadata["is_scanned_pdf"] is True
    assert mupdf_doc.metadata["warning_code"] == "SCANNED_PDF_DETECTED"
    assert mupdf_doc.metadata["scanned_pages_count"] == 1

    parsed = DocumentParser().parse(str(scanned_path))
    assert parsed.document.metadata["is_scanned_pdf"] is True
    assert any("SCANNED_PDF" in w for w in parsed.parser_warnings)
